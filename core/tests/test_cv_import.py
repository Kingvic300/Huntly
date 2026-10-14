from pathlib import Path

import pytest
from docx import Document as DocxDocument

from huntly.services.cv_import import ImportError_, extract_text, parse_cv
from huntly.services.documents import render_cv_docx, render_cv_pdf, tailor
from huntly.services.profile import Education, Experience, ProfileData

PROFILE = ProfileData(
    full_name="Ada Example",
    headline="Backend Engineer",
    email="ada@example.com",
    phone="+234 801 234 5678",
    location="Lagos, Nigeria",
    linkedin_url="https://linkedin.com/in/ada-example",
    github_url="https://github.com/adaex",
    summary="Backend engineer focused on reliable APIs.",
    skills=["Python", "FastAPI", "PostgreSQL", "Docker"],
    experience=[
        Experience(title="Backend Engineer", company="Acme Ltd", location="Lagos", start="Jan 2021", end="Present", highlights=["Built payment APIs in Python", "Moved deploys to Docker"]),
        Experience(title="Junior Developer", company="Beta Co", start="2019", end="2020", highlights=["Maintained Django apps"]),
    ],
    education=[Education(institution="University of Lagos", qualification="BSc Computer Science", start="2015", end="2019")],
    certifications=["AWS Certified Cloud Practitioner"],
    languages=["English", "Yoruba"],
)


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_round_trip_own_cv(tmp_path: Path, fmt):
    path = tmp_path / f"cv.{fmt}"
    t = tailor(PROFILE, "Backend Engineer", "Python")
    (render_cv_pdf if fmt == "pdf" else render_cv_docx)(t, path)
    r = parse_cv(extract_text(path, fmt))
    f = r["fields"]
    assert f["full_name"] == "Ada Example"
    assert f["email"] == "ada@example.com"
    assert f["location"] == "Lagos, Nigeria"
    assert f["linkedin_url"] == "https://linkedin.com/in/ada-example"
    assert f["github_url"] == "https://github.com/adaex"
    assert {"Python", "FastAPI", "PostgreSQL", "Docker"} <= set(f["skills"])
    titles = [e["title"] for e in f["experience"]]
    assert "Backend Engineer" in titles and "Junior Developer" in titles
    acme = next(e for e in f["experience"] if e["title"] == "Backend Engineer")
    assert acme["company"] == "Acme Ltd" and acme["location"] == "Lagos" and acme["start"] == "Jan 2021" and "Built payment APIs in Python" in acme["highlights"]
    assert any("University of Lagos" in e["institution"] for e in f["education"])
    assert "AWS Certified Cloud Practitioner" in f["certifications"]
    assert set(f["languages"]) == {"English", "Yoruba"}


def test_typical_handwritten_docx(tmp_path: Path):
    doc = DocxDocument()
    for line in [
        "JANE DOE",
        "Security Analyst",
        "jane.doe@mail.com | +44 7700 900123 | linkedin.com/in/janedoe | London, UK",
        "PROFESSIONAL SUMMARY",
        "Analyst with SOC experience.",
        "TECHNICAL SKILLS",
        "Languages: Python, Bash",
        "Tools: Splunk; Wireshark; Nmap",
        "WORK EXPERIENCE",
        "SOC Analyst at Contoso",
        "Mar 2022 – Present",
        "• Triaged alerts in Splunk",
        "• Wrote detection rules",
        "EDUCATION",
        "MSc Cyber Security",
        "University of Kent, 2021",
        "INTERESTS",
        "Chess",
    ]:
        doc.add_paragraph(line)
    path = tmp_path / "jane.docx"
    doc.save(path)
    f = parse_cv(extract_text(path, "docx"))["fields"]
    assert f["full_name"] == "Jane Doe" and f["headline"] == "Security Analyst"
    assert f["linkedin_url"] == "https://linkedin.com/in/janedoe"
    assert f["phone"].startswith("+44")
    assert f["location"] == "London, UK"
    assert {"Python", "Bash", "Splunk", "Wireshark", "Nmap"} <= set(f["skills"])
    assert f["experience"][0]["title"] == "SOC Analyst" and f["experience"][0]["company"] == "Contoso"
    assert f["experience"][0]["end"] == "Present" and len(f["experience"][0]["highlights"]) == 2
    assert "Chess" not in str(f)  # ignored section


def test_every_value_comes_from_the_document(tmp_path: Path):
    path = tmp_path / "cv.docx"
    render_cv_docx(tailor(PROFILE, "x", "y"), path)
    text = extract_text(path, "docx")
    f = parse_cv(text)["fields"]

    def strings(v):
        if isinstance(v, str):
            yield v
        elif isinstance(v, list):
            for x in v:
                yield from strings(x)
        elif isinstance(v, dict):
            for x in v.values():
                yield from strings(x)

    for value in strings(f):
        if value:
            assert value.replace("https://", "") in text or value in text or value.title() in text, value


def test_rejects_empty_or_image_only(tmp_path: Path):
    path = tmp_path / "empty.docx"
    DocxDocument().save(path)
    with pytest.raises(ImportError_):
        extract_text(path, "docx")


def test_parse_endpoint(client, tmp_path):
    import io

    path = tmp_path / "cv.pdf"
    render_cv_pdf(tailor(PROFILE, "x", "y"), path)
    doc = client.post("/api/documents", files={"file": ("cv.pdf", io.BytesIO(path.read_bytes()), "application/pdf")}).json()
    r = client.post(f"/api/documents/{doc['id']}/parse")
    assert r.status_code == 200 and r.json()["fields"]["full_name"] == "Ada Example"
    # Profile accepts the parsed links and normalises bare domains.
    prof = client.get("/api/profile").json()["profile"] | {"linkedin_url": "linkedin.com/in/x"}
    assert client.put("/api/profile", json=prof).json()["profile"]["linkedin_url"] == "https://linkedin.com/in/x"
    assert client.put("/api/profile", json=prof | {"linkedin_url": "not a url"}).status_code == 422
