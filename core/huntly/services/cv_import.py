"""Read an existing CV (PDF or DOCX) into profile fields, entirely on this machine.

Parsing is rule-based: every value returned is copied from the document's own text,
so nothing can be invented. The user reviews the result before it is saved.
"""

from __future__ import annotations

import re
from pathlib import Path

from huntly.security.sanitize import clean_text

MAX_PAGES = 12

HEADINGS: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "professional summary", "profile", "professional profile", "about me", "about", "objective", "career objective", "personal statement", "career summary"),
    "skills": ("skills", "technical skills", "core skills", "key skills", "competencies", "core competencies", "technologies", "tools", "skills and tools", "tech stack", "areas of expertise", "expertise"),
    "experience": ("experience", "work experience", "professional experience", "employment history", "work history", "career history", "employment", "relevant experience"),
    "education": ("education", "academic background", "education and training", "academic qualifications", "qualifications"),
    "projects": ("projects", "personal projects", "key projects", "selected projects", "open source"),
    "certifications": ("certifications", "certificates", "licenses and certifications", "licences and certifications", "certifications and licenses", "courses and certifications"),
    "languages": ("languages", "spoken languages"),
    "_ignore": ("references", "interests", "hobbies", "awards", "achievements", "volunteering", "volunteer experience", "publications", "referees"),
}
_HEADING_LOOKUP = {alias: key for key, aliases in HEADINGS.items() for alias in aliases}

EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
URL = re.compile(r"(?:https?://)?(?:www\.)?[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[a-z]{2,}(?:/[^\s,;|]*)?", re.I)
MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE = rf"(?:{MONTH}\s*\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
DATE_RANGE = re.compile(rf"(?P<start>{DATE})\s*(?:-|–|—|to)\s*(?P<end>{DATE}|present|current|now|date)", re.I)
BULLET = re.compile(r"^\s*(?:[•●▪◦‣∙·*\-–]|\d+[.)])\s+")
SPLIT_ITEMS = re.compile(r"\s*(?:,|;|\||•|·|●|▪|\t|\s{3,})\s*")
INSTITUTION = re.compile(r"(?i)\b(university|college|institute|school|polytechnic|academy|universit[éy])\b")


class ImportError_(Exception):
    pass


def extract_text(path: Path, fmt: str) -> str:
    if fmt == "pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise ImportError_("This PDF is password-protected. Export an unprotected copy and try again.")
            pages = [p.extract_text() or "" for p in reader.pages[:MAX_PAGES]]
        except PdfReadError as exc:
            raise ImportError_("This PDF couldn't be read") from exc
        text = "\n".join(pages)
    elif fmt == "docx":
        from docx import Document

        doc = Document(str(path))
        # Word list paragraphs carry their bullet in the style, not the text.
        parts = [("• " + p.text) if p.style is not None and "list" in (p.style.name or "").lower() and p.text.strip() else p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("  |  ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ImportError_("Only PDF and Word files can be imported")
    text = clean_text(text, 60_000)
    if len(text) < 40:
        raise ImportError_("No text found. If this CV is a scanned image, save it as a text-based PDF or Word file first.")
    return text


def _heading(line: str) -> str | None:
    key = re.sub(r"[^a-z& ]", "", line.lower()).replace("&", "and").strip()
    if len(key) > 40 or not key:
        return None
    return _HEADING_LOOKUP.get(key)


def _items(lines: list[str]) -> list[str]:
    out, seen = [], set()
    for line in lines:
        line = BULLET.sub("", line)
        if ":" in line and len(line.split(":", 1)[0]) < 30:
            line = line.split(":", 1)[1]
        for item in SPLIT_ITEMS.split(line):
            item = item.strip(" .")
            if 1 < len(item) <= 40 and item.lower() not in seen:
                seen.add(item.lower())
                out.append(item)
    return out


def _split_title(line: str) -> tuple[str, str]:
    for sep in (" at ", " @ ", " — ", " – ", " | ", " - "):
        if sep in line:
            a, b = line.split(sep, 1)
            return a.strip(" ,"), b.strip(" ,")
    if "," in line:
        a, b = line.split(",", 1)
        return a.strip(), b.strip()
    return line.strip(), ""


def _experience(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    cur: dict | None = None

    def new():
        nonlocal cur
        cur = {"title": "", "company": "", "location": "", "start": "", "end": "", "highlights": []}
        entries.append(cur)

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if BULLET.match(line):
            if cur is None:
                new()
            cur["highlights"].append(BULLET.sub("", line).strip())
            continue
        m = DATE_RANGE.search(line)
        rest = (line[: m.start()] + line[m.end() :]).strip(" |,–—-()") if m else line
        if m:
            if cur is None or cur["start"] or cur["highlights"]:
                new()
            cur["start"], cur["end"] = m.group("start"), m.group("end").title() if m.group("end").isalpha() else m.group("end")
            if rest and not cur["title"]:
                cur["title"], cur["company"] = _split_title(rest)
            elif rest and not cur["company"]:
                cur["company"] = rest
            elif rest and not cur["location"] and len(rest) <= 60:
                cur["location"] = rest
            continue
        if cur is not None and len(line) > 90:
            cur["highlights"].append(line)
            continue
        if cur is None or cur["highlights"] or (cur["title"] and cur["company"]):
            new()
        if not cur["title"]:
            cur["title"], company = _split_title(line)
            cur["company"] = cur["company"] or company
        elif not cur["company"]:
            cur["company"] = line
    return [e for e in entries if e["title"] or e["highlights"]][:20]


def _education(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    cur: dict | None = None
    for raw in lines:
        line = BULLET.sub("", raw).strip()
        if not line:
            continue
        m = DATE_RANGE.search(line) or re.search(r"\b(19|20)\d{2}\b", line)
        text = line
        if isinstance(m, re.Match) and m.re is DATE_RANGE:
            text = (line[: m.start()] + line[m.end() :]).strip(" |,–—-()")
        if cur is None or (INSTITUTION.search(text) and cur["institution"]) or (not INSTITUTION.search(text) and cur["qualification"] and text):
            cur = {"institution": "", "qualification": "", "start": "", "end": "", "details": ""}
            entries.append(cur)
        if isinstance(m, re.Match) and m.re is DATE_RANGE:
            cur["start"], cur["end"] = m.group("start"), m.group("end")
        elif m and not cur["end"]:
            cur["end"] = m.group(0)
            text = text.replace(m.group(0), "").strip(" |,–—-()")
        if not text:
            continue
        if INSTITUTION.search(text) and not cur["institution"]:
            cur["institution"] = text
        elif not cur["qualification"]:
            cur["qualification"] = text
        else:
            cur["details"] = (cur["details"] + " " + text).strip()
    return [e for e in entries if e["institution"] or e["qualification"]][:10]


def _projects(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        is_bullet = bool(BULLET.match(line))
        line = BULLET.sub("", line)
        if not is_bullet and len(line) <= 80 and (not entries or entries[-1]["description"]):
            name, desc = (line.split(":", 1) + [""])[:2] if ":" in line else (line, "")
            entries.append({"name": name.strip(), "description": desc.strip(), "url": "", "technologies": []})
        elif entries:
            entries[-1]["description"] = (entries[-1]["description"] + " " + line).strip()
        else:
            entries.append({"name": line[:80], "description": "", "url": "", "technologies": []})
    for e in entries:
        u = URL.search(e["description"]) or URL.search(e["name"])
        if u and "." in u.group(0) and ("/" in u.group(0) or u.group(0).startswith("http")):
            e["url"] = u.group(0) if u.group(0).startswith("http") else "https://" + u.group(0)
    return entries[:15]


def parse_cv(text: str) -> dict:
    lines = [l.strip() for l in text.splitlines()]
    sections: dict[str, list[str]] = {}
    header: list[str] = []
    current = None
    for line in lines:
        h = _heading(line) if line else None
        if h:
            current = h
            sections.setdefault(h, [])
            continue
        if current is None:
            header.append(line)
        else:
            sections[current].append(line)

    found: dict = {}
    top = "\n".join(header[:15]) or text[:1500]
    if m := EMAIL.search(top) or EMAIL.search(text):
        found["email"] = m.group(0)
    if m := PHONE.search(top):
        digits = re.sub(r"\D", "", m.group(0))
        if 8 <= len(digits) <= 15:
            found["phone"] = m.group(0).strip()
    links = []
    for u in URL.findall(text):
        if "@" in u or not ("/" in u or u.lower().startswith(("http", "www"))):
            continue
        full = u if u.lower().startswith("http") else "https://" + u
        low = full.lower()
        if "linkedin.com/" in low and "linkedin_url" not in found:
            found["linkedin_url"] = full.rstrip(".")
        elif "github.com/" in low and "github_url" not in found:
            found["github_url"] = full.rstrip(".")
        elif full not in links:
            links.append(full.rstrip("."))
    if links:
        found["portfolio_url"] = links[0]
        if len(links) > 1:
            found["links"] = links[1:5]

    for line in header[:6]:
        if not line or EMAIL.search(line) or URL.search(line) or PHONE.search(line):
            continue
        words = line.split()
        if "full_name" not in found and 2 <= len(words) <= 5 and all(re.match(r"^[A-Za-zÀ-ÿ'.-]+$", w) for w in words):
            found["full_name"] = " ".join(w if not w.isupper() else w.title() for w in words)
            continue
        if "full_name" in found and "headline" not in found and len(line) <= 90:
            found["headline"] = line
            break
    # Location: a short piece of a contact line that isn't an email, phone or link.
    for line in header[:8]:
        if "location" in found or line == found.get("headline") or line == found.get("full_name"):
            continue
        for part in re.split(r"\s*[|•·]\s*|\s{3,}", line):
            part = part.strip(" ,")
            if not part or EMAIL.search(part) or PHONE.search(part) or ("/" in part and URL.search(part)) or part.lower().startswith("www."):
                continue
            if re.fullmatch(r"[A-Za-zÀ-ÿ .'-]+,\s*[A-Za-zÀ-ÿ .'-]+", part) and len(part) <= 60:
                found["location"] = part
                break

    if sections.get("summary"):
        found["summary"] = " ".join(l for l in sections["summary"] if l).strip()[:2000]
    if sections.get("skills"):
        found["skills"] = _items([l for l in sections["skills"] if l])[:60]
    if sections.get("experience"):
        found["experience"] = _experience(sections["experience"])
    if sections.get("education"):
        found["education"] = _education(sections["education"])
    if sections.get("projects"):
        found["projects"] = _projects(sections["projects"])
    if sections.get("certifications"):
        found["certifications"] = [BULLET.sub("", l).strip() for l in sections["certifications"] if l.strip()][:20]
    if sections.get("languages"):
        found["languages"] = _items([l for l in sections["languages"] if l])[:10]

    warnings = []
    if not sections:
        warnings.append("No standard section headings were found, so only contact details could be read. Fill in the rest by hand.")
    for key in ("experience", "skills"):
        if key not in found:
            warnings.append(f"No {key} section was recognised.")
    return {"fields": found, "sections_found": sorted(k for k in sections if not k.startswith("_")), "warnings": warnings}
